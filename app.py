"""Two Degrees: Streamlit dashboard, validation and simulated batch replay."""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import time
from datetime import datetime, timedelta
from pathlib import Path

import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
APP_REVISION = "five-stage-incidents-v2"
GREEN, ORANGE, RED, NAVY, MUTED, TEAL = "#07876c", "#e88424", "#d83042", "#12324a", "#687c91", "#009bb2"
ICONS = {
    "milking": ":material/agriculture:",
    "initial_cooling": ":material/severe_cold:",
    "farm_storage": ":material/water_drop:",
    "loading": ":material/local_shipping:",
    "transport": ":material/local_shipping:",
    "delivery_storage": ":material/warehouse:",
}

INCIDENT_GUIDANCE = {
    "Missing readings": ["Check sensor power and cable", "Restore the sensor connection", "Confirm a fresh temperature reading"],
    "Sensor fault": ["Check the sensor diagnostics", "Use a spare sensor or manual probe", "Compare readings before resuming"],
    "High temperature": ["Check the cooling unit and doors", "Move milk to backup cooling", "Confirm a fresh reading between 0 and 4 °C"],
    "Low temperature": ["Check the cooler thermostat", "Reduce cooling and isolate the affected milk", "Confirm temperature is above 0 °C"],
    "Temperature boundary warning": ["Check whether temperature is rising or falling", "Adjust cooling before a limit is crossed", "Confirm temperature is between 0 and 4 °C"],
    "Cooling target missed": ["Check cooler flow and settings", "Use backup cooling and notify the operator", "Confirm the cooling target with a fresh reading"],
    "Power supply failure": ["Check the power supply status", "Switch to the backup power source", "Confirm power and cooling are restored"],
    "Refrigeration failure": ["Check refrigeration unit 1", "Switch to refrigeration unit 2", "Confirm the backup works and milk cools"],
    "Milk pump failure": ["Stop transfer and check the pump", "Use the backup pump or transfer route", "Confirm milk flow before continuing"],
    "Door left open": ["Close the tank or vehicle door", "Check the seal and cooling", "Confirm the door alarm clears"],
    "Loading delay": ["Keep the milk in chilled storage", "Arrange another loading slot or tanker", "Resume loading when the route is ready"],
    "Collection delay": ["Notify the collection operator", "Arrange a backup collection team", "Confirm collection resumes"],
    "Vehicle breakdown": ["Notify the dispatcher and keep cooling on", "Arrange a backup refrigerated vehicle", "Record transfer and confirm cooling"],
}


def dt(value):
    try:
        result = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if result.tzinfo is None:
            raise ValueError("Include a timezone offset in every timestamp.")
        return result
    except (TypeError, AttributeError, ValueError) as exc:
        raise ValueError("Invalid timestamp. Use ISO format, for example 2026-10-14T06:00:00+02:00.") from exc


def finite(value, label):
    if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value):
        raise ValueError(f"{label} must be a finite number.")
    return value


def text(value, label):
    if not isinstance(value, str) or not value.strip() or len(value) > 160:
        raise ValueError(f"{label} must be non-empty text of at most 160 characters.")
    return value


def adapt_legacy(data):
    """Accept the earlier shipment format without inventing upstream readings."""
    batches, readings = [], []
    for shipment in data.get("shipments", []):
        rows = [r for r in data.get("readings", []) if r.get("shipment_id") == shipment["shipment_id"]]
        if not rows:
            continue
        times = sorted(dt(r["timestamp"]) for r in rows)
        batches.append({
            "id": shipment["shipment_id"], "scenario": shipment.get("scenario", "Uploaded data"),
            "product": shipment.get("product", "Milk"), "quantity_litres": shipment.get("quantity_litres", 0),
            "owner": shipment.get("responsible_person", "Demo operator"),
            "origin": shipment.get("origin", "Origin"), "destination": shipment.get("destination", "Destination"),
            "stages": [{"id": "transport", "name": "Transport", "description": "Uploaded shipment data",
                        "start_at": times[0].isoformat(), "end_at": times[-1].isoformat(),
                        "min_c": shipment.get("temperature_min_c", 2), "max_c": shipment.get("temperature_max_c", 6),
                        "expected_interval_seconds": shipment.get("expected_reading_interval_seconds", 60),
                        "missing_after_seconds": shipment.get("missing_data_after_seconds", 120)}],
        })
        readings.extend({"batch_id": shipment["shipment_id"], "stage_id": "transport",
                          "timestamp": r["timestamp"], "temperature_c": r["temperature_c"],
                          "device_id": r.get("device_id", "SIM"), "is_simulated": True} for r in rows)
    return {"schema_version": "2.0", "metadata": {"source": "Legacy synthetic shipment data; transport only."},
            "batches": batches, "readings": readings}


def validate_data(raw):
    if not isinstance(raw, dict):
        raise ValueError("The JSON root must be an object.")
    if "shipments" in raw and "batches" not in raw:
        raw = adapt_legacy(raw)
    data = json.loads(json.dumps(raw))
    batches, readings = data.get("batches"), data.get("readings")
    if not isinstance(batches, list) or not 1 <= len(batches) <= 20:
        raise ValueError("Provide between 1 and 20 batches.")
    if not isinstance(readings, list) or not 1 <= len(readings) <= 10000:
        raise ValueError("Provide between 1 and 10,000 readings.")
    lookup = {}
    for batch in batches:
        if not isinstance(batch, dict):
            raise ValueError("Each batch must be an object.")
        bid = text(batch.get("id"), "Batch ID")
        if any(key[0] == bid for key in lookup):
            raise ValueError("Batch IDs must be unique.")
        for field in ("owner", "origin", "destination"):
            text(batch.get(field), field)
        if finite(batch.get("quantity_litres"), "Quantity") < 0:
            raise ValueError("Quantity cannot be negative.")
        stages = batch.get("stages")
        if not isinstance(stages, list) or not 1 <= len(stages) <= 12:
            raise ValueError("Each batch needs 1 to 12 stages.")
        previous_end = None
        for stage in stages:
            if not isinstance(stage, dict):
                raise ValueError("Each stage must be an object.")
            sid = text(stage.get("id"), "Stage ID")
            text(stage.get("name"), "Stage name")
            start, end = dt(stage.get("start_at")), dt(stage.get("end_at"))
            if start > end or (previous_end is not None and start <= previous_end):
                raise ValueError("Stages must be ordered, with non-overlapping time windows.")
            previous_end = end
            if (bid, sid) in lookup:
                raise ValueError("Stage IDs must be unique inside a batch.")
            for field in ("description", "location", "responsible_person", "demo_issue"):
                if field in stage:
                    text(stage[field], field)
            target = stage.get("cooling_target")
            if target is not None:
                if not isinstance(target, dict):
                    raise ValueError("cooling_target must be an object.")
                finite(target.get("max_c"), "Cooling target")
                if not start <= dt(target.get("due_at")) <= end:
                    raise ValueError("The cooling deadline must be inside its stage.")
            lower, upper = stage.get("min_c"), stage.get("max_c")
            if (lower is None) != (upper is None):
                raise ValueError("Provide both temperature limits, or set both to null.")
            if lower is not None and finite(lower, "Lower limit") >= finite(upper, "Upper limit"):
                raise ValueError("Lower limit must be less than upper limit.")
            interval = finite(stage.get("expected_interval_seconds", 60), "Reading interval")
            stale = finite(stage.get("missing_after_seconds", 120), "Missing-data timeout")
            if interval <= 0 or stale < interval:
                raise ValueError("Reading interval must be positive and timeout must be at least one interval.")
            stage.update(expected_interval_seconds=interval, missing_after_seconds=stale)
            lookup[(bid, sid)] = stage
    seen = set()
    for reading in readings:
        if not isinstance(reading, dict):
            raise ValueError("Each reading must be an object.")
        key = (text(reading.get("batch_id"), "Reading batch ID"),
               text(reading.get("stage_id"), "Reading stage ID"))
        if key not in lookup:
            raise ValueError("Every reading must reference an existing batch and stage.")
        timestamp = dt(reading.get("timestamp"))
        stage = lookup[key]
        if not dt(stage["start_at"]) <= timestamp <= dt(stage["end_at"]):
            raise ValueError("Reading timestamps must fall inside their stage window.")
        if not -100 <= finite(reading.get("temperature_c"), "Temperature") <= 150:
            raise ValueError("Temperature must be between -100 and 150 Celsius for this demo.")
        unique = (*key, timestamp)
        if unique in seen:
            raise ValueError("Duplicate timestamp within a batch stage.")
        seen.add(unique)
        reading["is_simulated"] = True
        if "device_id" in reading:
            text(reading["device_id"], "Device ID")
    readings.sort(key=lambda row: dt(row["timestamp"]))
    earliest = min(dt(stage["start_at"]) for batch in batches for stage in batch["stages"])
    latest = max(dt(stage["end_at"]) for batch in batches for stage in batch["stages"])
    if (latest - earliest).total_seconds() > 172800:
        raise ValueError("Use a sample covering at most 48 hours.")
    if not isinstance(data.get("metadata", {}), dict):
        raise ValueError("metadata must be an object.")
    events = data.get("scenario_events", [])
    if not isinstance(events, list):
        raise ValueError("scenario_events must be a list.")
    for event in events:
        if not isinstance(event, dict) or not event.get("batch_id") or not event.get("timestamp"):
            raise ValueError("Each scenario event needs a batch_id and timestamp.")
        dt(event["timestamp"])
        text(event["batch_id"], "Event batch ID")
        if event["batch_id"] not in {batch["id"] for batch in batches}:
            raise ValueError("Every event must reference an existing batch.")
        if "stage_id" in event:
            key = (event["batch_id"], text(event["stage_id"], "Event stage ID"))
            if key not in lookup:
                raise ValueError("Every event stage must exist in its batch.")
            stage = lookup[key]
            if not dt(stage["start_at"]) <= dt(event["timestamp"]) <= dt(stage["end_at"]):
                raise ValueError("Event timestamps must fall inside their stage window.")
        text(event.get("type"), "Event type")
        if "description" in event:
            text(event["description"], "Event description")
        if event["type"] in ("fault_started", "fault_resolved"):
            if "stage_id" not in event:
                raise ValueError("Fault events must reference a stage.")
            text(event.get("fault_id"), "Fault ID")
            if event["type"] == "fault_started":
                if event.get("incident_type") not in INCIDENT_GUIDANCE:
                    raise ValueError("Unsupported incident_type in fault event.")
                if event.get("severity", "critical") not in ("warning", "critical"):
                    raise ValueError("Fault severity must be warning or critical.")
    faults = {}
    for event in sorted(events, key=lambda item: dt(item["timestamp"])):
        if event["type"] not in ("fault_started", "fault_resolved"):
            continue
        key = (event["batch_id"], event["stage_id"], event["fault_id"])
        if event["type"] == "fault_started":
            if key in faults:
                raise ValueError("Fault IDs must be unique inside a stage.")
            faults[key] = [dt(event["timestamp"]), False]
        elif key not in faults or faults[key][1] or dt(event["timestamp"]) <= faults[key][0]:
            raise ValueError("A fault recovery must follow its matching fault start.")
        else:
            faults[key][1] = True
    # Keep the journey and its exports consistent without changing the source JSON.
    for batch in data["batches"]:
        batch["stages"] = [stage for stage in batch["stages"] if stage["id"] != "delivery_storage"]
        if not batch["stages"]:
            raise ValueError("Each batch needs a stage before delivery storage.")
        for stage in batch["stages"]:
            if stage.get("min_c") is not None:
                stage.update(min_c=0, max_c=4)
            if stage.get("cooling_target"):
                stage["cooling_target"]["max_c"] = 4
    retained = {(batch["id"], stage["id"]) for batch in data["batches"] for stage in batch["stages"]}
    data["readings"] = [row for row in data["readings"] if (row["batch_id"], row["stage_id"]) in retained]
    batch_ends = {batch["id"]: max(dt(stage["end_at"]) for stage in batch["stages"])
                  for batch in data["batches"]}
    data["scenario_events"] = [event for event in events
                               if event["batch_id"] in batch_ends
                               and event.get("stage_id") != "delivery_storage"
                               and dt(event["timestamp"]) <= batch_ends[event["batch_id"]]]
    metadata = data.setdefault("metadata", {})
    playback = metadata.setdefault("playback", {})
    if not isinstance(playback, dict):
        raise ValueError("metadata.playback must be an object.")
    interval = finite(playback.get("sample_interval_seconds", 60), "Playback interval")
    if interval < 1:
        raise ValueError("Playback interval must be at least one second.")
    speed = finite(playback.get("recommended_seconds_per_step", 1), "Playback speed")
    if speed <= 0:
        raise ValueError("Playback speed must be positive.")
    playback["recommended_seconds_per_step"] = min([.5, 1., 2., 3., 5.], key=lambda value: abs(value - speed))
    playback["start_at"] = min(dt(stage["start_at"]) for batch in data["batches"]
                               for stage in batch["stages"]).isoformat()
    playback["end_at"] = max(batch_ends.values()).isoformat()
    metadata["threshold_note"] = (
        "Demo chilled-stage colours: 0 and 4 °C orange; between 0 and 4 °C green; "
        "below 0 or above 4 °C red. Collection and cooling progress remain observational."
    )
    data["scenario_events"].sort(key=lambda event: dt(event["timestamp"]))
    return data


def timeline(data):
    metadata = data.get("metadata", {})
    playback = metadata.get("playback", {})
    starts = [dt(stage["start_at"]) for batch in data["batches"] for stage in batch["stages"]]
    ends = [dt(stage["end_at"]) for batch in data["batches"] for stage in batch["stages"]]
    tick = max(1, int(playback.get("sample_interval_seconds", min(
        stage["expected_interval_seconds"] for batch in data["batches"] for stage in batch["stages"]
    ))))
    start, end = dt(playback.get("start_at", min(starts).isoformat())), dt(playback.get("end_at", max(ends).isoformat()))
    count = int((end - start).total_seconds() / tick) + 1
    if count > 20000:
        raise ValueError("Too many playback ticks.")
    return sorted(set([start, end] + [start + timedelta(seconds=tick * i) for i in range(count)] +
                      [dt(row["timestamp"]) for row in data["readings"]] +
                      starts + ends))


def stage_rows(data, batch_id, stage_id, now):
    return [row for row in data["readings"] if row["batch_id"] == batch_id and row["stage_id"] == stage_id
            and dt(row["timestamp"]) <= now]


def assessment_limits(stage, timestamp=None):
    """Warm collection/cooling is observational until the cooling deadline."""
    target = stage.get("cooling_target")
    if target:
        if timestamp is None or dt(timestamp) < dt(target["due_at"]):
            return None
        return 0, 4
    return (0, 4) if stage.get("min_c") is not None else None


def risk_level(stage, temperature, timestamp=None):
    limits = assessment_limits(stage, timestamp)
    if limits is None:
        return "observational"
    lower, upper = limits
    if temperature < lower or temperature > upper:
        return "critical"
    if temperature == lower or temperature == upper:
        return "warning"
    return "ideal"


def reading_status(stage, temperature, timestamp=None):
    level = risk_level(stage, temperature, timestamp)
    return "unassessed" if level == "observational" else "in_range" if level == "ideal" else "out_of_range"


def alert_flag(stage, temperature, timestamp=None):
    return int(risk_level(stage, temperature, timestamp) in ("warning", "critical"))


def latest_fresh_reading(stage, rows, now):
    effective = min(now, dt(stage["end_at"]))
    last = next((row for row in reversed(rows) if dt(row["timestamp"]) <= effective), None)
    if last and (effective - dt(last["timestamp"])).total_seconds() < stage["missing_after_seconds"]:
        return last
    return None


def state_for(stage, rows, now):
    start = dt(stage["start_at"])
    if now < start:
        return "waiting"
    last = latest_fresh_reading(stage, rows, now)
    if last is None:
        return "data_unavailable"
    return reading_status(stage, last["temperature_c"], min(now, dt(stage["end_at"])).isoformat())


def cooling_endpoint_status(stage, rows, now):
    target = stage.get("cooling_target")
    if not target or now < dt(target["due_at"]):
        return "pending"
    fresh = latest_fresh_reading(stage, rows, dt(target["due_at"]))
    if fresh is None:
        return "unknown"
    return "within_target" if 0 <= fresh["temperature_c"] <= target["max_c"] else "outside_target"


def incidents_for(data, batch, stage, now):
    """Keep temperature incidents open across sensor loss; require observed recovery."""
    start, end = dt(stage["start_at"]), min(now, dt(stage["end_at"]))
    if end < start:
        return []
    rows = stage_rows(data, batch["id"], stage["id"], end)
    times = {start, end}
    first_timeout = start + timedelta(seconds=stage["missing_after_seconds"])
    if first_timeout <= end:
        times.add(first_timeout)
    for row in rows:
        moment = dt(row["timestamp"])
        times.add(moment)
        expiry = moment + timedelta(seconds=stage["missing_after_seconds"])
        if expiry <= end:
            times.add(expiry)
    target = stage.get("cooling_target")
    if target and dt(target["due_at"]) <= end:
        times.add(dt(target["due_at"]))
    episodes, opened = [], {}

    def begin(kind, moment, severity, description):
        if kind not in opened:
            opened[kind] = {
                "id": f'{batch["id"]}/{stage["id"]}/{kind}/{moment.isoformat()}',
                "batch_id": batch["id"], "stage_id": stage["id"], "stage": stage["name"],
                "type": kind, "severity": severity, "source": "Temperature / sensor readings",
                "description": description, "started_at": moment.isoformat(), "ended_at": None,
            }

    def recover(kind, moment):
        if kind in opened:
            incident = opened.pop(kind)
            incident["ended_at"] = moment.isoformat()
            episodes.append(incident)

    pointer, last = 0, None
    thermal = ("High temperature", "Low temperature", "Temperature boundary warning", "Cooling target missed")
    for moment in sorted(times):
        new_reading = False
        while pointer < len(rows) and dt(rows[pointer]["timestamp"]) <= moment:
            last = rows[pointer]
            pointer += 1
            new_reading = True
        age = (moment - (dt(last["timestamp"]) if last else start)).total_seconds()
        fresh = last is not None and age < stage["missing_after_seconds"]
        if not fresh and age >= stage["missing_after_seconds"]:
            begin("Missing readings", moment, "warning", "No fresh temperature sample received before the sensor timeout.")
        if fresh:
            recover("Missing readings", moment)
        # A stale value never closes a temperature incident or proves recovery.
        if not fresh or not (new_reading or target and moment == dt(target["due_at"])):
            continue
        temperature = last["temperature_c"]
        level = risk_level(stage, temperature, moment.isoformat())
        if level == "ideal":
            for kind in thermal:
                recover(kind, moment)
        elif target and moment >= dt(target["due_at"]) and level == "critical":
            begin("Cooling target missed", moment, "critical", "Milk did not reach the 0–4 °C cooling target by its deadline.")
        elif not target and level in ("warning", "critical"):
            kind = "Temperature boundary warning" if level == "warning" else "Low temperature" if temperature < 0 else "High temperature"
            begin(kind, moment, level, f"Measured {temperature:g} °C; chilled demo range is 0–4 °C with orange boundaries.")
    episodes.extend(opened.values())
    # Equipment causes are supplied explicitly by the simulator, never inferred from temperature alone.
    faults = {}
    for event in data.get("scenario_events", []):
        if event["batch_id"] != batch["id"] or event.get("stage_id") != stage["id"] or dt(event["timestamp"]) > end:
            continue
        if event["type"] == "fault_started":
            faults[event["fault_id"]] = {
                "id": f'{batch["id"]}/{stage["id"]}/fault/{event["fault_id"]}',
                "batch_id": batch["id"], "stage_id": stage["id"], "stage": stage["name"],
                "type": event["incident_type"], "severity": event.get("severity", "critical"),
                "source": "Simulated equipment event", "description": event.get("description", event["incident_type"]),
                "started_at": event["timestamp"], "ended_at": None,
            }
        elif event["type"] == "fault_resolved" and event["fault_id"] in faults:
            faults[event["fault_id"]]["ended_at"] = event["timestamp"]
    episodes.extend(faults.values())
    return sorted(episodes, key=lambda incident: (dt(incident["started_at"]), incident["id"]))


def all_incidents(data, batch, now):
    return sorted([incident for stage in batch["stages"] for incident in incidents_for(data, batch, stage, now)],
                  key=lambda incident: (dt(incident["started_at"]), incident["id"]))


def graph_for(stage, rows):
    fig = go.Figure()
    for i in range(1, len(rows)):
        previous, row = rows[i - 1], rows[i]
        if (dt(row["timestamp"]) - dt(previous["timestamp"])).total_seconds() > stage["expected_interval_seconds"] * 1.5:
            continue
        levels = {risk_level(stage, previous["temperature_c"], previous["timestamp"]), risk_level(stage, row["temperature_c"], row["timestamp"])}
        color = RED if "critical" in levels else ORANGE if "warning" in levels else GREEN if stage.get("min_c") is not None else NAVY
        fig.add_trace(go.Scatter(x=[dt(previous["timestamp"]), dt(row["timestamp"])],
                                 y=[previous["temperature_c"], row["temperature_c"]], mode="lines",
                                 line=dict(color=color, width=3), showlegend=False, hoverinfo="skip"))
    groups = [("ideal", "Within ideal range", GREEN), ("warning", "Warning: at boundary", ORANGE),
              ("critical", "Critical", RED), ("observational", "Cooling / collection", NAVY)]
    for level, label, color in groups:
        points = [row for row in rows if risk_level(stage, row["temperature_c"], row["timestamp"]) == level]
        if points:
            fig.add_trace(go.Scatter(x=[dt(row["timestamp"]) for row in points],
                                     y=[row["temperature_c"] for row in points], mode="markers", name=label,
                                     marker=dict(color=color, size=8),
                                     hovertemplate="%{x|%H:%M:%S}<br>%{y:.1f}°C<extra>" + label + "</extra>"))
    if stage.get("min_c") is not None:
        fig.add_hrect(y0=stage["min_c"], y1=stage["max_c"], fillcolor=GREEN, opacity=.05, line_width=0)
        for limit in (stage["min_c"], stage["max_c"]):
            fig.add_hline(y=limit, line_dash="dash", line_color="#9aafc1",
                          annotation_text=f"{limit:g}°C limit", annotation_position="top left")
    elif stage.get("cooling_target"):
        fig.add_hline(y=stage["cooling_target"]["max_c"], line_dash="dot", line_color="#6d8ca3",
                      annotation_text=f'Endpoint target {stage["cooling_target"]["max_c"]:g}°C')
    fig.update_layout(height=285, margin=dict(l=55, r=18, t=24, b=12), paper_bgcolor="white",
                      plot_bgcolor="white", font=dict(family="Arial", color=NAVY, size=12),
                      legend=dict(orientation="h", y=-.22, x=0), xaxis=dict(tickformat="%H:%M", gridcolor="#eef2f6", range=[dt(stage["start_at"]), dt(stage["end_at"])], nticks=6),
                      yaxis=dict(title="Temperature (°C)", automargin=True, gridcolor="#e9eef3", zeroline=False),
                      hovermode="closest")
    return fig


def csv_data(rows, stage):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["timestamp", "temperature_c", "status", "alert_flag", "batch_id", "stage_id", "is_simulated"])
    for row in rows:
        writer.writerow([row["timestamp"], row["temperature_c"], risk_level(stage, row["temperature_c"], row["timestamp"]),
                         alert_flag(stage, row["temperature_c"], row["timestamp"]),
                         row["batch_id"], row["stage_id"], True])
    return output.getvalue()


def incident_action(actions, incident_id, action_type):
    return next((action for action in reversed(actions)
                 if action.get("incident_id") == incident_id and action["type"] == action_type), None)


def incident_status(incident, actions):
    if incident["ended_at"] and incident_action(actions, incident["id"], "Incident report completed"):
        return "Report complete"
    if incident["ended_at"]:
        return "Recovered"
    if incident_action(actions, incident["id"], "Alert acknowledged"):
        return "Acknowledged"
    return "Active alert"


def records_payload(data, batch, now, incidents, actions, events, run_id):
    stage_map = {stage["id"]: stage for stage in batch["stages"]}
    readings = []
    for row in data["readings"]:
        if row["batch_id"] != batch["id"] or dt(row["timestamp"]) > now:
            continue
        stage = stage_map[row["stage_id"]]
        enriched = dict(row)
        enriched["status"] = risk_level(stage, row["temperature_c"], row["timestamp"])
        enriched["alert_flag"] = alert_flag(stage, row["temperature_c"], row["timestamp"])
        readings.append(enriched)
    alert_rows = []
    for number, incident in enumerate(incidents, 1):
        alert_rows.append({
            "alert_id": f"ALT-{number:03d}", "batch_id": batch["id"], "issue": incident["type"],
            "stage": incident["stage"], "severity": incident["severity"], "source": incident["source"],
            "description": incident["description"], "incident_id": incident["id"], "started_at": incident["started_at"],
            "recovered_at": incident["ended_at"], "status": incident_status(incident, actions),
            "alert_flag": 0 if incident["ended_at"] else 1,
        })
    return {"run_id": run_id, "as_of": now.isoformat(), "is_simulated": True, "batch": batch,
            "readings": readings, "alerts_and_incidents": alert_rows,
            "scenario_events": events, "operator_records": actions}


def excel_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def excel_records(payload):
    """Create a readable multi-sheet workbook entirely in memory."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Batch summary"
    batch = payload["batch"]
    summary_rows = [
        ("Two Degrees record export", "SIMULATED DATA"), ("Exported as of", payload["as_of"]),
        ("Batch ID", batch["id"]), ("Scenario", batch.get("scenario", "")),
        ("Product", batch.get("product", "Milk")), ("Quantity (L)", batch["quantity_litres"]),
        ("Origin", batch["origin"]), ("Destination", batch["destination"]),
        ("Vehicle", batch.get("vehicle_id", "")), ("Responsible operator", batch["owner"]),
        ("Readings exported", len(payload["readings"])),
        ("Alerts and incidents", len(payload["alerts_and_incidents"])),
    ]
    for row in summary_rows:
        summary.append([excel_value(value) for value in row])
    summary["A1"].font = Font(bold=True, color="FFFFFF")
    summary["B1"].font = Font(bold=True, color="FFFFFF")
    summary["A1"].fill = summary["B1"].fill = PatternFill("solid", fgColor="009BB2")

    sheets = [
        ("Readings", payload["readings"]), ("Alerts and incidents", payload["alerts_and_incidents"]),
        ("Journey events", payload["scenario_events"]), ("Operator records", payload["operator_records"]),
    ]
    for title, records in sheets:
        sheet = workbook.create_sheet(title[:31])
        if not records:
            sheet.append(["No records at this replay time"])
            continue
        headers = list(dict.fromkeys(key for record in records for key in record))
        sheet.append(headers)
        for record in records:
            sheet.append([excel_value(record.get(header)) for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12324A")
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(headers) + 1):
            width = min(38, max(12, max(len(str(sheet.cell(row, column).value or ""))
                                       for row in range(1, sheet.max_row + 1)) + 2))
            sheet.column_dimensions[get_column_letter(column)].width = width
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 60
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


CSS = """<style>
:root{color-scheme:light}
.stApp{background:radial-gradient(circle at 88% 2%,#def5f7 0,transparent 34%),#f4f8fc;color:#12324a}
[data-testid='stHeader']{background:transparent}.block-container{max-width:1660px;padding:2.2rem 2rem 3rem}
[data-testid='stVerticalBlock']{gap:.72rem}h1,h2,h3{color:#12324a;letter-spacing:-.035em}
h1{font-size:1.85rem!important;padding:0!important}h3{font-size:1.08rem!important;padding:.1rem 0 .3rem!important}
button{border-radius:10px!important;transition:.15s ease}button:focus-visible{outline:3px solid #64c6d7!important;outline-offset:3px}
[data-testid='stBaseButton-primary']{background:linear-gradient(110deg,#009bb2,#2878d4)!important;border:0!important;color:white!important}
[data-testid='stBaseButton-primary'] p{color:white!important}[data-baseweb='select']>div{background:#fff;border-color:#d9e5ed;border-radius:10px}
[data-testid='stFileUploaderDropzone']{background:#fff;border:1px dashed #9cc5d3;border-radius:12px}
.brand{display:flex;align-items:center;gap:20px;min-height:166px}.brand img{flex-shrink:0;width:156px;height:156px;object-fit:contain}
.brand strong{font-size:25px;color:#092f4e;letter-spacing:-.8px}.brand small{display:block;font-size:10px;font-weight:700;letter-spacing:2px;color:#1494a7;margin-top:2px}
.st-key-nav{border-bottom:1px solid #dce7ee;padding-bottom:10px;margin-bottom:3px}.st-key-nav button{min-height:39px}
.eyebrow{font-size:10px;font-weight:700;letter-spacing:1.6px;color:#598393;margin-bottom:5px}.titleline{display:flex;gap:12px;align-items:center;flex-wrap:wrap}
.pill{font-size:10px;font-weight:700;letter-spacing:.7px;background:#dff2f4;color:#087e90;border-radius:20px;padding:5px 9px}.summary{display:flex;gap:28px;align-items:center;flex-wrap:wrap;font-size:11px;color:#6c8293;padding:3px 0}
.summary b{display:block;color:#12324a;font-size:13px;font-weight:600;margin-top:2px}.clock{font-size:12px;color:#58788c;text-align:right}.clock b{color:#12324a;font-size:17px;font-variant-numeric:tabular-nums}
.st-key-replay{background:#ffffffc9;border:1px solid #dce7ef;border-radius:14px;padding:9px 14px}.sectionline{display:flex;align-items:center;justify-content:space-between;margin-top:4px}.sectionline strong{font-size:15px}.sectionline span{font-size:11px;color:#718699}
div[class*='st-key-card_']{position:relative!important;border:1px solid #dae5ed;border-radius:15px;background:linear-gradient(145deg,#fff,#f3f9fc);box-shadow:0 5px 13px #17364b0a;overflow:visible;gap:0!important;transform-style:preserve-3d;transition:transform .2s ease,box-shadow .2s ease,border-color .2s ease}
div[class*='st-key-card_']:hover{transform:translateY(-6px) rotateX(2deg);border-color:#56b9ca;box-shadow:0 14px 25px #17364b1a;z-index:4}
div[class*='st-key-card_'] [data-testid='stElementContainer']:has([data-testid='stButton']){position:static!important;height:0!important;min-height:0!important}div[class*='st-key-card_'] [data-testid='stButton']{position:static!important;height:0!important}
div[class*='st-key-card_'] [data-testid='stButton'] button{position:absolute!important;inset:0!important;width:100%!important;height:100%!important;z-index:3;background:transparent!important;border:0!important;color:transparent!important;box-shadow:none!important}div[class*='st-key-card_'] [data-testid='stButton'] button p{opacity:0}
.cardbody{height:236px;padding:13px 12px;box-sizing:border-box;pointer-events:none;position:relative}.cardtop{display:flex;align-items:center;gap:7px}.stagetitle{font-size:12px;font-weight:700;color:#234259;line-height:1.15}.stagephase{margin-left:auto;font-size:9px;font-weight:700;color:#718697;background:#eef3f6;padding:3px 6px;border-radius:6px}
.process{display:flex;align-items:center;gap:9px;margin:10px 0 5px}.tank{flex-shrink:0;width:62px;height:108px;border:2px solid #84b5c5;border-radius:8px 8px 10px 10px;position:relative;overflow:hidden;background:#f6fbfd}.liquid{position:absolute;bottom:0;left:0;right:0;background:linear-gradient(#42c0d1,#0d91aa);transition:height .5s ease}.tank svg{position:absolute;z-index:1;width:36px;height:36px;left:11px;top:34px;color:#fff;filter:drop-shadow(0 1px 1px #0b6680)}
.cardvalue{font-size:22px;font-weight:700;letter-spacing:-.7px;font-variant-numeric:tabular-nums}.cardvalue small{font-size:13px;font-weight:500}.cardstatus{font-size:11px;font-weight:600;min-height:30px;line-height:1.3;margin-top:4px;overflow-wrap:anywhere}.progress{height:12px;border-radius:3px;background:#e5edf2;margin-top:8px;overflow:hidden}.progress span{display:block;height:100%;background:linear-gradient(90deg,#10a5b9,#2878d4)}.cardfoot{font-size:9px;color:#7a8fa0;display:flex;justify-content:space-between;margin-top:6px}
.good{color:#07876c}.warning{color:#d36b08}.bad{color:#d83042}.neutral{color:#678297}.phase-active{color:#067f91;background:#dff4f6}.phase-complete{color:#08735d;background:#e1f4ec}.phase-warning{color:#a95108;background:#fff0de}
.st-key-chart_panel,.st-key-log_panel,.st-key-record_panel,.st-key-alert_register,.st-key-alert_detail,.st-key-demo_panel{background:#fff;border:1px solid #dce6ee!important;border-radius:16px!important;padding:17px!important;box-shadow:0 5px 18px #17364b06}.st-key-incident_workspace{max-width:1200px;margin:0 auto}
.mini{display:grid;grid-template-columns:1fr 1.3fr 1fr;gap:12px;margin:8px 0 0;padding:10px 0;border-top:1px solid #edf1f5}.mini span{display:block;color:#7a8c9a;font-size:9px;text-transform:uppercase;letter-spacing:.7px}.mini b{display:block;font-size:17px;font-weight:600;color:#26485e;margin-top:3px}
.reading-table{width:100%;border-collapse:collapse;font-size:11px}.reading-table th{background:#f3f7fa;color:#6b8191;text-align:left;font-weight:600;padding:9px 7px}.reading-table td{padding:9px 7px;border-bottom:1px solid #edf2f6;font-variant-numeric:tabular-nums}.empty{height:150px;display:flex;align-items:center;justify-content:center;color:#8194a3;font-size:12px}
.alert-kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:8px 0 16px}.kpi{background:#fff;border:1px solid #dce6ee;border-radius:14px;padding:13px 15px;box-shadow:0 5px 15px #17364b05}.kpi span{font-size:10px;color:#718697;text-transform:uppercase;letter-spacing:.7px}.kpi b{display:block;font-size:25px;margin-top:4px}.kpi small{font-size:10px;color:#8194a3}.alert-detail-head{display:flex;align-items:flex-start;justify-content:space-between;gap:12px}.alert-id{font-size:10px;color:#7290a0;letter-spacing:1px}.alert-title{font-size:21px;font-weight:700;color:#12324a}.status-chip{font-size:10px;font-weight:700;padding:5px 8px;border-radius:12px;background:#fff0de;color:#a95108}.detail-grid{display:grid;grid-template-columns:repeat(2,1fr);gap:12px;margin:16px 0}.detail-grid div{background:#f6f9fb;border-radius:10px;padding:10px}.detail-grid span{font-size:9px;color:#8194a3;text-transform:uppercase;letter-spacing:.6px}.detail-grid b{display:block;font-size:12px;margin-top:3px}.steps{display:grid;grid-template-columns:repeat(3,1fr);gap:8px;margin:12px 0}.step{background:#eef8fa;border-radius:9px;padding:10px;font-size:10px;color:#406276}.step b{display:block;color:#008aa0;font-size:12px;margin-bottom:3px}
.overview-alerts{display:grid;grid-template-columns:repeat(3,1fr);gap:10px;margin:10px 0}.overview-alert{background:#fff;border:1px solid #dce6ee;border-radius:12px;padding:11px 13px}.overview-alert span{font-size:10px;color:#718697}.overview-alert b{display:block;font-size:20px;margin-top:2px}.footer{font-size:10px;color:#8b9ca9;padding-top:10px}
@media(max-width:1000px){.cardbody{height:232px}.stagetitle{font-size:10px}.cardvalue{font-size:19px}.tank{width:38px;height:76px}.tank svg{width:24px;height:24px;left:5px;top:24px}}
@media(max-width:750px){.block-container{padding:2.7rem .8rem 2rem}.brand{min-height:126px}.brand img{width:116px;height:116px}.brand strong{font-size:21px}.brand small{display:none}.clock{text-align:left}.titleline h1{font-size:1.55rem!important}.summary{gap:15px}.alert-kpis,.overview-alerts{grid-template-columns:1fr 1fr}.detail-grid,.steps{grid-template-columns:1fr}.mini b{font-size:14px}
.st-key-nav [data-testid='stHorizontalBlock'],.st-key-replay [data-testid='stHorizontalBlock']{flex-wrap:wrap!important;gap:7px!important}.st-key-nav [data-testid='stHorizontalBlock']>[data-testid='stColumn']{min-width:0!important;flex:1 1 32%!important;width:auto!important}.st-key-nav [data-testid='stHorizontalBlock']>[data-testid='stColumn']:first-child{flex:1 1 100%!important}.st-key-replay [data-testid='stHorizontalBlock']>[data-testid='stColumn']{min-width:0!important;flex:1 1 27%!important;width:auto!important}.st-key-replay [data-testid='stHorizontalBlock']>[data-testid='stColumn']:nth-child(4){flex:1 1 100%!important}}
</style>"""


def esc(value):
    return html.escape(str(value))


def load_dataset(raw):
    data = validate_data(raw)
    ticks = timeline(data)
    playback = data.get('metadata', {}).get('playback', {})
    default = data.get('metadata', {}).get('default_batch_id')
    bid = default if default in [b['id'] for b in data['batches']] else data['batches'][0]['id']
    batch = next(b for b in data['batches'] if b['id'] == bid)
    ss = st.session_state
    ss.update(app_revision=APP_REVISION, data=data, ticks=ticks, index=0, batch_id=bid, batch_choice=bid,
              selected_stage=batch['stages'][0]['id'], follow_active=True, follow_control=True, playing=False,
              speed=float(playback.get('recommended_seconds_per_step', 1)),
              last_tick=time.monotonic(), actions=[], run_id=str(time.time_ns()))


def active_stage(batch, now):
    return next((s for s in batch['stages'] if dt(s['start_at']) <= now <= dt(s['end_at'])), None)


def select_stage(sid):
    st.session_state.selected_stage = sid
    st.session_state.follow_active = False
    st.session_state.follow_control = False


def change_follow():
    st.session_state.follow_active = st.session_state.follow_control


def jump_to_incident(timestamp, stage_id):
    ss = st.session_state
    ss.index = next(index for index, tick in enumerate(ss.ticks) if tick >= dt(timestamp))
    ss.playing = False
    ss.selected_stage = stage_id
    ss.follow_active = False
    ss.follow_control = False
    ss.last_tick = time.monotonic()


def change_batch():
    ss = st.session_state
    ss.batch_id = ss.batch_choice
    batch = next(b for b in ss.data['batches'] if b['id'] == ss.batch_id)
    current = active_stage(batch, ss.ticks[ss.index])
    ss.selected_stage = current['id'] if current else batch['stages'][0]['id']
    ss.follow_active = True
    ss.follow_control = True


SVG_PATHS = {
    'milking':'<path d="M4 21V8l8-5 8 5v13H4Z M9 21v-7h6v7 M8 9h.01 M16 9h.01"/>',
    'initial_cooling':'<path d="M12 2v20 M3.3 7l17.4 10 M3.3 17 20.7 7 M9 4l3 3 3-3 M9 20l3-3 3 3"/>',
    'farm_storage':'<ellipse cx="12" cy="5" rx="7" ry="3"/><path d="M5 5v14c0 4 14 4 14 0V5 M5 12c0 4 14 4 14 0"/>',
    'loading':'<path d="M3 8h10v11H3z M13 12h5l3 4v3h-8 M8 3v8 M5 8l3 3 3-3"/><circle cx="6" cy="20" r="2"/><circle cx="18" cy="20" r="2"/>',
    'transport':'<path d="M2 5h12v13H2z M14 10h4l4 5v3h-8"/><circle cx="6" cy="19" r="2"/><circle cx="18" cy="19" r="2"/>',
    'delivery_storage':'<path d="M3 21V8l9-5 9 5v13H3Z M7 21V11h10v10 M7 15h10 M7 18h10"/>'}


def render_cards(data, batch, now, current):
    ss = st.session_state
    st.markdown('<div class="sectionline"><strong>Batch journey</strong><span>Click a stage to inspect</span></div>', unsafe_allow_html=True)
    columns = st.columns(len(batch['stages']), gap='small')
    styles = []
    for index, stage in enumerate(batch['stages']):
        rows = stage_rows(data, batch['id'], stage['id'], now)
        status = state_for(stage, rows, now)
        completed = now >= dt(stage['end_at'])
        delivered = stage['id'] == 'transport' and completed
        phase = 'Delivered' if delivered else 'Completed' if completed else 'Active' if current and current['id'] == stage['id'] else 'Upcoming'
        stage_incidents = incidents_for(data, batch, stage, now)
        active = [incident for incident in stage_incidents if incident['ended_at'] is None]
        last = latest_fresh_reading(stage, rows, now)
        level = risk_level(stage, last['temperature_c'], min(now, dt(stage['end_at'])).isoformat()) if last else None
        color = {'critical': 'bad', 'warning': 'warning', 'ideal': 'good'}.get(level, 'neutral')
        if status == 'data_unavailable':
            color = 'warning'
        severity = 'critical' if any(incident['severity'] == 'critical' for incident in active) else 'warning' if active else None
        phase_class = 'phase-complete' if completed else 'phase-active' if phase == 'Active' else ''
        value = f'{last["temperature_c"]:.1f}<small>°C</small>' if last else 'Waiting' if status == 'waiting' else 'Unknown'
        if delivered:
            label = f'Delivered · {len(active)} active alerts' if active else 'Batch delivered'
        elif active:
            label = active[-1]['type']
        elif stage_incidents:
            label = f'{len(stage_incidents)} incident(s) recovered'
        elif completed:
            label = 'Completed'
        else:
            label = {'waiting': 'Waiting for stage', 'in_range': 'Within ideal range',
                     'data_unavailable': 'Waiting for sensor', 'unassessed': 'Cooling in progress' if stage.get('cooling_target') else 'Collection in progress',
                     'out_of_range': 'Temperature alert'}[status]
        if stage.get('cooling_target') and now >= dt(stage['cooling_target']['due_at']) and not stage_incidents:
            label = {'within_target': 'Cooling target reached', 'outside_target': 'Cooling target missed',
                     'unknown': 'Cooling result unknown', 'pending': 'Cooling in progress'}[cooling_endpoint_status(stage, rows, now)]
        start, end = dt(stage['start_at']), dt(stage['end_at'])
        progress = 0 if now < start else 100 if now >= end else round((now - start).total_seconds() / (end - start).total_seconds() * 100)
        selected = ss.selected_stage == stage['id']
        if selected:
            styles.append(f'.stApp .st-key-card_{index}{{border-color:#009bb2;box-shadow:0 0 0 1px #009bb2;background:#effbfc}}')
        if severity:
            border, background = ('#d83042', '#fff2f3') if severity == 'critical' else ('#e88424', '#fff5e9')
            styles.append(f'.stApp .st-key-card_{index}{{border-color:{border};background:{background}}}')
        icon = SVG_PATHS.get(stage['id'], SVG_PATHS['farm_storage'])
        tooltip = esc(stage.get('description', stage['name']))
        body = (f'<div class="cardbody" title="{tooltip}"><div class="cardtop"><span class="stagetitle">{esc(stage["name"])}</span>'
                f'<span class="stagephase {phase_class}">{phase}</span></div>'
                f'<div class="process"><div class="tank"><div class="liquid" style="height:{progress}%"></div>'
                f'<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.6">{icon}</svg></div>'
                f'<div><div class="cardvalue {color}">{value}</div><div class="cardstatus {"bad" if severity == "critical" else "warning" if severity else "neutral"}">{esc(label)}</div></div></div>'
                f'<div class="progress"><span style="width:{progress}%"></span></div>'
                f'<div class="cardfoot"><span>{progress}%</span><span>{"Viewing" if selected else "Inspect ↗"}</span></div></div>')
        with columns[index], st.container(key=f'card_{index}'):
            st.markdown(body, unsafe_allow_html=True)
            st.button(f'View {stage["name"]} history', key=f'choose_{index}', width='stretch',
                      on_click=select_stage, args=(stage['id'],))
    st.markdown('<style>'+''.join(styles)+'</style>', unsafe_allow_html=True)


def table_html(rows, stage):
    result = '<table class="reading-table"><thead><tr><th>Time</th><th>Temperature</th><th>Status</th></tr></thead><tbody>'
    for r in rows:
        level = risk_level(stage, r['temperature_c'], r['timestamp'])
        color = {'ideal': 'good', 'warning': 'warning', 'critical': 'bad', 'observational': 'neutral'}[level]
        label = {'ideal': 'Within ideal range', 'warning': 'Boundary warning', 'critical': 'Critical',
                 'observational': 'Cooling / collection'}[level]
        result += f'<tr class="{color}"><td>{dt(r["timestamp"]):%H:%M}</td><td>{r["temperature_c"]:.1f} °C</td><td>{label}</td></tr>'
    return result+'</tbody></table>'


def alert_metrics(incidents, actions):
    active = sum(1 for incident in incidents if not incident['ended_at'])
    acknowledged = sum(1 for incident in incidents
                       if incident_action(actions, incident['id'], 'Alert acknowledged'))
    recovered = sum(1 for incident in incidents if incident['ended_at'])
    pending = sum(1 for incident in incidents
                  if not incident_action(actions, incident['id'], 'Incident report completed'))
    return active, acknowledged, recovered, pending


def render_alert_kpis(incidents, actions, compact=False):
    active, acknowledged, recovered, pending = alert_metrics(incidents, actions)
    if compact:
        st.markdown(
            '<div class="overview-alerts">'
            f'<div class="overview-alert"><span>Active alerts</span><b class="{"warning" if active else "good"}">{active}</b></div>'
            f'<div class="overview-alert"><span>Recovered incidents</span><b class="good">{recovered}</b></div>'
            f'<div class="overview-alert"><span>Reports pending</span><b class="{"warning" if pending else "good"}">{pending}</b></div>'
            '</div>', unsafe_allow_html=True)
        return
    cards = [('Active alerts', active, 'Needs attention'), ('Acknowledged', acknowledged, 'Seen by operator'),
             ('Recovered', recovered, 'Reading returned'), ('Reports pending', pending, 'Close the record')]
    body = '<div class="alert-kpis">'
    for label, value, helper in cards:
        tone = 'warning' if value and label in ('Active alerts', 'Reports pending') else 'good'
        body += f'<div class="kpi"><span>{label}</span><b class="{tone}">{value}</b><small>{helper}</small></div>'
    st.markdown(body + '</div>', unsafe_allow_html=True)


def render_incident_page(data, batch, incidents, actions, now):
    render_alert_kpis(incidents, actions)
    if not incidents:
        with st.container(key='alert_detail'):
            st.markdown('### No alerts yet')
            st.caption('Track this batch. System-generated alerts will appear here when a reading leaves the ideal range or data becomes unavailable.')
        return

    options = {f'ALT-{number:03d} | {incident["type"]} | {incident["stage"]}': incident
               for number, incident in enumerate(incidents, 1)}
    register, detail = st.columns([.78, 1.42], gap='medium')
    with register, st.container(key='alert_register'):
        st.subheader('Alert register')
        choice = st.selectbox('Choose an alert', list(reversed(options)), label_visibility='collapsed')
        selected = options[choice]
        rows = []
        for label, incident in reversed(list(options.items())):
            rows.append({'Alert': label.split(' | ')[0], 'Issue': incident['type'],
                         'Time': dt(incident['started_at']).strftime('%H:%M'),
                         'Status': incident_status(incident, actions)})
        listing = '<table class="reading-table"><thead><tr><th>Alert</th><th>Time</th><th>Status</th></tr></thead><tbody>'
        for row in rows:
            tone = 'warning' if row['Status'] in ('Active alert', 'Acknowledged') else 'good'
            listing += (f'<tr><td><b>{esc(row["Alert"])}</b><br><small>{esc(row["Issue"])}</small></td>'
                        f'<td>{esc(row["Time"])}</td><td class="{tone}">{esc(row["Status"])}</td></tr>')
        st.markdown(listing + '</tbody></table>', unsafe_allow_html=True)

    with detail, st.container(key='alert_detail'):
        selected_number = list(options.values()).index(selected) + 1
        alert_id = f'ALT-{selected_number:03d}'
        status = incident_status(selected, actions)
        stage = next(stage for stage in batch['stages'] if stage['id'] == selected['stage_id'])
        rows = stage_rows(data, batch['id'], stage['id'], now)
        fresh = latest_fresh_reading(stage, rows, now)
        latest = f'{fresh["temperature_c"]:.1f} °C' if fresh else 'Unknown (no fresh reading)'
        flag = 0 if selected['ended_at'] else 1
        st.markdown(
            f'<div class="alert-detail-head"><div><div class="alert-id">{alert_id}</div>'
            f'<div class="alert-title">{esc(selected["type"])}</div></div>'
            f'<span class="status-chip">{esc(status)}</span></div>'
            f'<div class="detail-grid"><div><span>Batch</span><b>{esc(batch["id"])}</b></div>'
            f'<div><span>Alert flag</span><b>{flag}</b></div>'
            f'<div><span>Stage</span><b>{esc(selected["stage"])}</b></div>'
            f'<div><span>Location</span><b>{esc(stage.get("location", "Unknown"))}</b></div>'
            f'<div><span>Detected</span><b>{dt(selected["started_at"]):%d %b, %H:%M}</b></div>'
            f'<div><span>Latest reading</span><b>{latest}</b></div>'
            f'<div><span>Container / vehicle</span><b>{esc(batch.get("container_id") or batch.get("vehicle_id", "Unassigned"))}</b></div>'
            f'<div><span>Operator</span><b>{esc(batch["owner"])}</b></div></div>', unsafe_allow_html=True)
        st.caption(f'{selected["severity"].title()} · {selected["source"]}')
        st.write(selected['description'])
        guidance = INCIDENT_GUIDANCE.get(selected['type'], INCIDENT_GUIDANCE['High temperature'])
        st.markdown('<div class="steps">' + ''.join(
            f'<div class="step"><b>Step {index}</b>{esc(item)}</div>'
            for index, item in enumerate(guidance, 1)) + '</div>', unsafe_allow_html=True)
        acknowledged = bool(incident_action(actions, selected['id'], 'Alert acknowledged'))
        complete = bool(incident_action(actions, selected['id'], 'Incident report completed'))
        first, second = st.columns(2)
        if first.button('Acknowledged' if acknowledged else 'Acknowledge alert',
                        disabled=acknowledged, type='primary' if not acknowledged else 'secondary',
                        width='stretch', key=f'ack_{selected_number}'):
            st.session_state.actions.append(dict(timestamp=now.isoformat(), batch_id=batch['id'],
                incident_id=selected['id'], actor=batch['owner'], type='Alert acknowledged',
                note='Operator confirmed the system alert.'))
            st.rerun()
        if second.button('Report complete' if complete else 'Complete incident report',
                         disabled=complete or not acknowledged or selected['ended_at'] is None, width='stretch', key=f'close_{selected_number}'):
            st.session_state.actions.append(dict(timestamp=now.isoformat(), batch_id=batch['id'],
                incident_id=selected['id'], actor=batch['owner'], type='Incident report completed',
                note='Incident details reviewed and record completed.'))
            st.rerun()
        if not acknowledged:
            st.caption('Acknowledge the system alert before completing its report.')
        elif selected['ended_at'] is None:
            st.caption('Recovery must be confirmed before the incident report can be completed.')


def render_demo_page():
    ss = st.session_state
    with st.container(key='demo_panel'):
        st.subheader('Data')
        st.caption('Upload another simulated journey or reset this proof of concept to the bundled sample.')
        st.caption('Five stages: milking → initial cooling → farm storage → loading → transport / delivered.')
        upload = st.file_uploader('Upload journey JSON', type=['json'], max_upload_size=2, key='upload')
        if upload is not None:
            fingerprint = hashlib.sha256(upload.getvalue()).hexdigest()
            if fingerprint != ss.get('upload_fingerprint'):
                try:
                    load_dataset(json.loads(upload.getvalue().decode('utf-8-sig')))
                    ss.upload_fingerprint = fingerprint
                    st.success('Data loaded. Choose Overview, then Track batch.')
                except (ValueError, KeyError, TypeError) as exc:
                    st.error(f'Cannot load this file: {exc}')
        reset, sample = st.columns(2)
        if reset.button('Load bundled sample', width='stretch', type='primary'):
            load_dataset(json.loads((ROOT/'sample_data.json').read_text(encoding='utf-8')))
            if upload is not None:
                ss.upload_fingerprint = hashlib.sha256(upload.getvalue()).hexdigest()
            st.rerun()
        sample.download_button('Download current dataset', json.dumps(ss.data, indent=2),
                               'sample_data.json', 'application/json', width='stretch')
        ss.speed = st.select_slider('Playback speed: seconds per simulated minute',
                                    options=[.5, 1., 2., 3., 5.], value=ss.speed)
        st.markdown(f'**{len(ss.data["batches"])} batches**  ·  **{len(ss.data["readings"])} readings**  ·  No physical sensor required')

        st.subheader('Demo scenario guide')
        st.caption('Planned scenarios only. Alerts appear when replay reaches their detection time.')
        guide = []
        for batch in ss.data['batches']:
            item = {'Batch': batch['id']}
            for stage in batch['stages']:
                item[stage['name']] = stage.get('demo_issue', 'See journey readings')
            guide.append(item)
        st.dataframe(guide, hide_index=True, width='stretch')


def main():
    st.set_page_config(page_title='Two Degrees | Milk journey',page_icon='🥛',layout='wide',initial_sidebar_state='collapsed')
    st.markdown(CSS,unsafe_allow_html=True)
    ss=st.session_state
    if 'data' not in ss or ss.get('app_revision') != APP_REVISION: load_dataset(json.loads((ROOT/'sample_data.json').read_text(encoding='utf-8')))
    ss.setdefault('follow_active', True)
    ss.setdefault('follow_control', ss.follow_active)
    with st.container(key='nav'):
        brand,*links=st.columns([4,1,1,1,1.05],vertical_alignment='center',gap='small')
        with brand:
            logo=ROOT/'logo-mark-transparent.png'
            img=f'<img alt="Two Degrees logo" src="data:image/png;base64,{__import__("base64").b64encode(logo.read_bytes()).decode()}">' if logo.exists() else ''
            st.markdown(f'<div class="brand">{img}<div><strong>Two Degrees</strong><small>COLD-CHAIN RESILIENCE</small></div></div>',unsafe_allow_html=True)
        for col,page in zip(links,['Overview','Incidents','Records','Data']):
            if col.button(page,key='nav_'+page.lower(),type='primary' if ss.get('page','Overview')==page else 'secondary',width='stretch'):
                ss.page=page; st.rerun()

    if ss.get('page') == 'Data':
        render_demo_page()
        st.markdown('<div class="footer">Two Degrees · Proof of concept · Simulated data only.</div>', unsafe_allow_html=True)
        return

    @st.fragment(run_every=.5 if ss.playing else None)
    def dashboard():
        if ss.playing and time.monotonic()-ss.last_tick>=ss.speed:
            ss.index=min(ss.index+1,len(ss.ticks)-1); ss.last_tick=time.monotonic()
            if ss.index==len(ss.ticks)-1:
                ss.playing=False; st.rerun()
        now=ss.ticks[ss.index]; data=ss.data
        batches={b['id']:b for b in data['batches']}; batch=batches[ss.batch_id]
        if 'batch_choice' not in ss:
            ss.batch_choice = ss.batch_id
        current=active_stage(batch,now)
        if ss.follow_active and current: ss.selected_stage=current['id']
        title,clock=st.columns([3,1],vertical_alignment='center')
        with title:
            st.markdown('<div class="eyebrow">FROM FARM TO DELIVERY</div><div class="titleline"><h1>Milk journey monitoring</h1><span class="pill">SIMULATED DATA</span></div>',unsafe_allow_html=True)
        with clock:
            st.markdown(f'<div class="clock">{now:%d %b %Y}<br><b>{now:%H:%M}</b> · {"Playing" if ss.playing else "Finished" if ss.index==len(ss.ticks)-1 else "Paused"}</div>',unsafe_allow_html=True)
        pick,details=st.columns([1,2],vertical_alignment='center',gap='medium')
        with pick:
            st.selectbox('Batch',list(batches),format_func=lambda bid:f'{bid} · {batches[bid].get("scenario","Uploaded data")}',key='batch_choice',on_change=change_batch,label_visibility='collapsed')
        with details:
            st.markdown(f'<div class="summary"><div>Volume<b>{batch["quantity_litres"]:g} L</b></div><div>Route<b>{esc(batch["origin"].replace(" (Demo)",""))} → {esc(batch["destination"].replace(" (Demo)",""))}</b></div><div>Vehicle<b>{esc(batch.get("vehicle_id","Unassigned"))}</b></div></div>',unsafe_allow_html=True)
        with st.container(key='replay'):
            play,nextcol,restart,seek,follow=st.columns([1,.95,.8,2.5,1.2],gap='small',vertical_alignment='center')
            if play.button('Pause tracking' if ss.playing else 'Track batch',type='primary',width='stretch',disabled=ss.index==len(ss.ticks)-1):
                ss.playing=not ss.playing; ss.last_tick=time.monotonic(); st.rerun()
            if nextcol.button('Next update',width='stretch',disabled=ss.playing or ss.index==len(ss.ticks)-1):
                ss.index+=1; st.rerun()
            if restart.button('Restart',width='stretch'):
                # Callback-safe reset without changing the already-rendered batch widget.
                ss.index=0; ss.playing=False; ss.actions=[]; ss.run_id=str(time.time_ns()); ss.follow_active=True; ss.follow_control=True
                ss.selected_stage=batch['stages'][0]['id']; ss.last_tick=time.monotonic(); st.rerun()
            with seek:
                requested=st.slider('Replay position',0,len(ss.ticks)-1,ss.index,key=f'seek_{ss.run_id}_{ss.index}',disabled=ss.playing,label_visibility='collapsed')
                if requested!=ss.index: ss.index=requested; ss.last_tick=time.monotonic(); st.rerun()
            with follow: st.toggle('Follow journey',key='follow_control',on_change=change_follow)
        planned = all_incidents(data, batch, ss.ticks[-1])
        upcoming = next((incident for incident in planned if dt(incident['started_at']) > now), None)
        jump, hint = st.columns([1, 4], vertical_alignment='center')
        jump.button('Next incident', disabled=upcoming is None, width='stretch',
                    on_click=jump_to_incident,
                    args=(upcoming['started_at'], upcoming['stage_id']) if upcoming else (None, None))
        hint.caption('Demo shortcut: jump to the next incident for the selected batch.')
        incidents=all_incidents(data,batch,now)
        actions=[a for a in ss.actions if a['batch_id']==batch['id'] and dt(a['timestamp'])<=now]
        events=sorted([e for e in data.get('scenario_events',[]) if e['batch_id']==batch['id'] and dt(e['timestamp'])<=now],key=lambda e:dt(e['timestamp']))
        page=ss.get('page','Overview')
        if page=='Records':
            with st.container(key='record_panel'):
                st.subheader(f'{batch["id"]} · Records')
                payload=records_payload(data,batch,now,incidents,actions,events,ss.run_id)
                st.caption(f'{len(payload["readings"])} readings · {len(incidents)} alerts and incidents · {len(actions)} operator records')
                excel,json_col=st.columns(2)
                excel.download_button('Export Excel',excel_records(payload),f'{batch["id"]}-records.xlsx',
                                      'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                      type='primary',width='stretch')
                json_col.download_button('Export JSON',json.dumps(payload,indent=2),f'{batch["id"]}-records.json',
                                         'application/json',width='stretch')
                st.caption('Excel is the readable report. JSON keeps the same information for technical use.')
            return
        if page=='Overview':
            render_cards(data,batch,now,current)
            render_alert_kpis(incidents,actions,compact=True)
            stage=next(s for s in batch['stages'] if s['id']==ss.selected_stage)
            rows=stage_rows(data,batch['id'],stage['id'],now); status=state_for(stage,rows,now)
            left,right=st.columns([1.7,1],gap='medium')
            with left,st.container(key='chart_panel'):
                st.subheader(f'{stage["name"]} · Temperature history')
                st.caption(f'{stage.get("location","Stage readings")} · {stage.get("responsible_person",batch["owner"])}')
                if stage.get('min_c') is not None:
                    st.caption('Green: 0 < temperature < 4 °C · Orange: exactly 0 or 4 °C · Red: below 0 or above 4 °C')
                latest=f'{rows[-1]["temperature_c"]:.1f} °C' if rows and status!='data_unavailable' else 'Unknown' if status=='data_unavailable' else 'Not started'
                limits=f'{stage["min_c"]:g}–{stage["max_c"]:g} °C' if stage.get('min_c') is not None else 'Before cooling'
                if stage.get('cooling_target'): limits=f'≤ {stage["cooling_target"]["max_c"]:g} °C at endpoint'
                stamp=dt(rows[-1]['timestamp']).strftime('%H:%M') if rows else 'No sample'
                st.markdown(f'<div class="mini"><div><span>Latest temperature</span><b>{latest}</b></div><div><span>Demo assessment</span><b>{limits}</b></div><div><span>Last sample</span><b>{stamp}</b></div></div>',unsafe_allow_html=True)
                if rows: st.plotly_chart(graph_for(stage,rows),width='stretch',theme=None,key='history_chart',config={'displayModeBar':False})
                else: st.markdown('<div class="empty">Readings will appear when this stage begins.</div>',unsafe_allow_html=True)
                if status=='data_unavailable': st.warning('Current temperature is unknown. The graph preserves the gap in readings.')
            with right,st.container(key='log_panel'):
                st.subheader('Reading log')
                st.caption(f'{batch["id"]} · {stage["name"]} · Latest 6 readings')
                st.markdown(table_html(rows[-6:],stage),unsafe_allow_html=True)
                if not rows: st.caption('No readings yet.')
                with st.expander(f'Full history ({len(rows)} readings)'):
                    st.markdown('<div style="max-height:320px;overflow:auto">'+table_html(rows,stage)+'</div>',unsafe_allow_html=True)
                st.download_button('Download stage log',csv_data(rows,stage),f'{batch["id"]}-{stage["id"]}.csv','text/csv',width='stretch')
        if page=='Incidents':
            st.markdown('<div class="eyebrow">SYSTEM ALERTS AND OPERATOR RESPONSE</div>', unsafe_allow_html=True)
            st.markdown('### Incident centre')
            with st.container(key='incident_workspace'):
                render_incident_page(data,batch,incidents,actions,now)
        st.markdown('<div class="footer">Two Degrees · Proof of concept · Demo thresholds only. Temperature status does not certify product safety.</div>',unsafe_allow_html=True)
    dashboard()


if __name__=='__main__':
    main()
